"""Render an account ``Statement`` as a branded PDF (same design language as the certificate
and the installment schedule) or as an Excel workbook (three sheets: movements, summary by
type, holdings). Both renderers only lay out what ``statement_service`` computed."""

from __future__ import annotations

import datetime as dt
import decimal
import io

from reportlab.pdfgen import canvas

from app.services.certificate_service import (
    _CX,
    _GOLD,
    _GREEN,
    _GREEN_D,
    _H,
    _INK,
    _MUTED,
    _PANEL,
    _W,
    _clip,
    _divider,
    _draw_brand,
    _fit_centred,
    _wrap,
)
from app.services.certificate_service import _spaced_centred as _spaced_centred_raw
from app.services.installment_pdf import _RED, _page_frame
from app.services.statement_service import Statement

_BOTTOM = 112.0
_ROW_H = 15.0
_C_DATE = 62.0
_C_TYPE = 122.0
_C_DESC = 222.0
_C_STATUS = 372.0
_C_AMOUNT = 470.0  # right-aligned
_C_BAL = 536.0  # right-aligned


def _spaced_centred(c: canvas.Canvas, *args) -> None:
    """The shared helper sets PDF character spacing, which persists in the graphics state and
    would widen every later line (columns then overlap). Contain it."""
    c.saveState()
    _spaced_centred_raw(c, *args)
    c.restoreState()


def _fmt(v: decimal.Decimal, signed: bool = False) -> str:
    text = f"{abs(v):,.2f}"
    if not signed:
        return f"-{text}" if v < 0 else text
    return f"+{text}" if v >= 0 else f"-{text}"


def _day(d: dt.date | dt.datetime) -> str:
    return d.strftime("%d %b %Y")


def _period(stmt: Statement) -> str:
    return f"{_day(stmt.start)} – {_day(stmt.end)}"


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def _footer(c: canvas.Canvas, stmt: Statement, page: int) -> None:
    c.setFillColor(_MUTED)
    c.setFont("Times-Roman", 8.5)
    note = (
        "Generated from the Capimax PropShare wallet ledger. Times are UTC; the period includes "
        "both dates. Amounts in progress (for example a withdrawal being paid) are shown as "
        "they were booked. This statement is informational and is not a tax document."
    )
    fy = 96
    for ln in _wrap(note, "Times-Roman", 8.5, 400):
        c.drawCentredString(_CX, fy, ln)
        fy -= 11
    c.setFillColor(_GOLD)
    c.setFont("Helvetica", 8)
    c.drawCentredString(_CX, 62, f"{stmt.ref}   •   page {page}   •   capimaxpropshare.com")


def _movements_header(c: canvas.Canvas, y: float) -> float:
    c.setFillColor(_MUTED)
    c.setFont("Helvetica-Bold", 7.5)
    c.drawString(_C_DATE, y, "DATE")
    c.drawString(_C_TYPE, y, "TYPE")
    c.drawString(_C_DESC, y, "DESCRIPTION")
    c.drawString(_C_STATUS, y, "STATUS")
    c.drawRightString(_C_AMOUNT, y, "AMOUNT")
    c.drawRightString(_C_BAL, y, "BALANCE")
    c.setStrokeColor(_GOLD)
    c.setLineWidth(0.8)
    c.line(_C_DATE, y - 6, _C_BAL, y - 6)
    return y - 19


class _Pdf:
    """A canvas that knows how to break pages with a continued-title header."""

    def __init__(self, stmt: Statement) -> None:
        self.stmt = stmt
        self.buf = io.BytesIO()
        self.c = canvas.Canvas(self.buf, pagesize=(_W, _H), pageCompression=0)
        self.c.setTitle("Capimax PropShare — Account Statement")
        self.page = 1
        _page_frame(self.c)

    def new_page(self, title: str) -> float:
        _footer(self.c, self.stmt, self.page)
        self.c.showPage()
        self.page += 1
        _page_frame(self.c)
        _spaced_centred(self.c, _CX, 780, "Times-Bold", 15, title, 1.0, _GREEN_D)
        return 748.0

    def finish(self) -> bytes:
        _footer(self.c, self.stmt, self.page)
        self.c.showPage()
        self.c.save()
        return self.buf.getvalue()


def render_pdf(stmt: Statement) -> bytes:
    pdf = _Pdf(stmt)
    c = pdf.c

    _draw_brand(c, _CX)
    _divider(c, _CX, 712, 150)
    _spaced_centred(c, _CX, 676, "Times-Bold", 27, "Account Statement", 1.0, _GREEN_D)
    c.setFillColor(_MUTED)
    c.setFont("Times-Italic", 11.5)
    c.drawCentredString(_CX, 658, f"{_period(stmt)}  (UTC)")
    _fit_centred(c, _CX, 630, "Times-Bold", 15, stmt.holder, 440, _GREEN)
    c.setFillColor(_INK)
    c.setFont("Times-Roman", 10.5)
    c.drawCentredString(_CX, 613, f"{stmt.email}   •   {stmt.ref}")

    # summary panel: 3 columns x 2 rows
    px, pw, ph, py = 56.0, _W - 112.0, 84.0, 510.0
    c.setFillColor(_PANEL)
    c.setStrokeColor(_GOLD)
    c.setLineWidth(0.8)
    c.roundRect(px, py, pw, ph, 6, stroke=1, fill=1)
    cur = stmt.currency
    fields = [
        ("OPENING BALANCE", f"{_fmt(stmt.opening)} {cur}"),
        ("MONEY IN", f"{_fmt(stmt.money_in)} {cur}"),
        ("MONEY OUT", f"{_fmt(stmt.money_out)} {cur}"),
        ("CLOSING BALANCE", f"{_fmt(stmt.closing)} {cur}"),
        ("MOVEMENTS", str(len(stmt.rows))),
        ("ISSUED", f"{_day(stmt.generated_at)} {stmt.generated_at:%H:%M} UTC"),
    ]
    col_w = pw / 3
    for i, (label, val) in enumerate(fields):
        col, row = i % 3, i // 3
        x = px + 18 + col * col_w
        y = py + ph - 24 - row * 34
        c.setFillColor(_MUTED)
        c.setFont("Helvetica", 6.8)
        c.drawString(x, y, label)
        strong = label in ("CLOSING BALANCE",)
        c.setFillColor(_GREEN_D if strong else _INK)
        c.setFont("Times-Bold" if strong else "Times-Roman", 10.5)
        c.drawString(x, y - 15, _clip(val, "Times-Roman", 10.5, col_w - 24))

    # movements
    y = py - 26
    c.setFillColor(_GREEN_D)
    c.setFont("Times-Bold", 12)
    c.drawString(_C_DATE, y, "Wallet movements")
    y = _movements_header(c, y - 20)
    if not stmt.rows:
        c.setFillColor(_MUTED)
        c.setFont("Times-Italic", 10)
        c.drawString(_C_DATE, y, "No wallet movements in this period.")
        y -= _ROW_H
    for i, r in enumerate(stmt.rows):
        if y < _BOTTOM:
            y = _movements_header(c, pdf.new_page("Wallet movements (continued)"))
        if i % 2:
            c.setFillColor(_PANEL)
            c.rect(_C_DATE - 6, y - 4, _C_BAL - _C_DATE + 12, _ROW_H, stroke=0, fill=1)
        c.setFillColor(_MUTED)
        c.setFont("Times-Roman", 8.5)
        c.drawString(_C_DATE, y, _day(r.at))
        c.setFillColor(_INK)
        c.drawString(_C_TYPE, y, _clip(r.type_label, "Times-Roman", 8.5, _C_DESC - _C_TYPE - 6))
        c.drawString(
            _C_DESC, y, _clip(r.description or "-", "Times-Roman", 8.5, _C_STATUS - _C_DESC - 8)
        )
        c.setFillColor(_MUTED)
        c.drawString(_C_STATUS, y, r.status)
        c.setFillColor(_GREEN if r.amount >= 0 else _RED)
        c.setFont("Times-Bold", 8.5)
        c.drawRightString(_C_AMOUNT, y, _fmt(r.amount, signed=True))
        c.setFillColor(_INK)
        c.setFont("Times-Roman", 8.5)
        c.drawRightString(_C_BAL, y, _fmt(r.balance))
        y -= _ROW_H

    # summary by type + holdings
    def section(title: str, head: tuple[str, str, str], lines: list[tuple[str, str, str]]):
        nonlocal y
        if y < _BOTTOM + 60:  # title + header + one row must fit, else start a page
            y = pdf.new_page(title)
        else:
            y -= 18
            c.setFillColor(_GREEN_D)
            c.setFont("Times-Bold", 12)
            c.drawString(_C_DATE, y, title)
            y -= 20
        c.setFillColor(_MUTED)
        c.setFont("Helvetica-Bold", 7.5)
        c.drawString(_C_DATE, y, head[0])
        c.drawRightString(_C_AMOUNT, y, head[1])
        c.drawRightString(_C_BAL, y, head[2])
        c.setStrokeColor(_GOLD)
        c.line(_C_DATE, y - 6, _C_BAL, y - 6)
        y -= 19
        for a, b, d in lines or [("None", "", "")]:
            if y < _BOTTOM:
                y = pdf.new_page(f"{title} (continued)")
            c.setFillColor(_INK)
            c.setFont("Times-Roman", 9)
            c.drawString(_C_DATE, y, _clip(a, "Times-Roman", 9, _C_AMOUNT - _C_DATE - 80))
            c.drawRightString(_C_AMOUNT, y, b)
            c.drawRightString(_C_BAL, y, d)
            y -= _ROW_H

    section(
        "Summary by type",
        ("TYPE", "MONEY IN", "MONEY OUT"),
        [(t.label, _fmt(t.money_in), _fmt(t.money_out)) for t in stmt.totals],
    )
    section(
        f"Holdings on {_day(stmt.end)}",
        ("PROPERTY", "", "UNITS"),
        [(h.property_title, "", str(h.units)) for h in stmt.holdings],
    )
    return pdf.finish()


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def render_xlsx(stmt: Statement) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    money = "#,##0.00;[Red]-#,##0.00"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="0F6E4C")
    head_font = Font(bold=True, color="FFFFFF")

    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    meta = [
        ("Capimax PropShare — Account statement", None),
        ("Holder", stmt.holder),
        ("Email", stmt.email),
        ("Reference", stmt.ref),
        ("Period (UTC, inclusive)", f"{stmt.start.isoformat()} to {stmt.end.isoformat()}"),
        ("Currency", stmt.currency),
        ("Opening balance", stmt.opening),
        ("Money in", stmt.money_in),
        ("Money out", stmt.money_out),
        ("Closing balance", stmt.closing),
        ("Issued (UTC)", stmt.generated_at.strftime("%Y-%m-%d %H:%M")),
    ]
    for label, value in meta:
        ws.append([label, value])
    ws["A1"].font = Font(bold=True, size=14, color="0F6E4C")
    for row in ws.iter_rows(min_row=2, max_row=len(meta), max_col=1):
        row[0].font = bold
    for r in range(7, 11):
        ws.cell(row=r, column=2).number_format = money
    ws.append([])

    header = ["Date (UTC)", "Type", "Description", "Status", "Amount", "Balance", "Reference"]
    ws.append(header)
    head_row = ws.max_row
    for cell in ws[head_row]:
        cell.font, cell.fill = head_font, head_fill
    for r in stmt.rows:
        ws.append(
            [
                r.at.astimezone(dt.UTC).replace(tzinfo=None),
                r.type_label,
                r.description,
                r.status,
                r.amount,
                r.balance,
                r.reference,
            ]
        )
        ws.cell(row=ws.max_row, column=1).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(row=ws.max_row, column=5).number_format = money
        ws.cell(row=ws.max_row, column=6).number_format = money
    if not stmt.rows:
        ws.append(["No wallet movements in this period."])
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    for col, width in zip("ABCDEFG", (24, 24, 46, 13, 15, 15, 12), strict=True):
        ws.column_dimensions[col].width = width

    totals = wb.create_sheet("Summary by type")
    totals.append(["Type", "Money in", "Money out"])
    for cell in totals[1]:
        cell.font, cell.fill = head_font, head_fill
    for t in stmt.totals:
        totals.append([t.label, t.money_in, t.money_out])
    totals.append(["Total", stmt.money_in, stmt.money_out])
    for cell in totals[totals.max_row]:
        cell.font = bold
    for row in totals.iter_rows(min_row=2, min_col=2, max_col=3):
        for cell in row:
            cell.number_format = money
    totals.column_dimensions["A"].width = 28
    totals.column_dimensions["B"].width = 16
    totals.column_dimensions["C"].width = 16

    holdings = wb.create_sheet("Holdings")
    holdings.append([f"Property (units held on {stmt.end.isoformat()})", "Units"])
    for cell in holdings[1]:
        cell.font, cell.fill = head_font, head_fill
    for h in stmt.holdings:
        holdings.append([h.property_title, h.units])
    holdings.column_dimensions["A"].width = 48
    holdings.column_dimensions["B"].width = 10
    for sheet in (ws, totals, holdings):
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
