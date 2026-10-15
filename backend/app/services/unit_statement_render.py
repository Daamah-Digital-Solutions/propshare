"""Render a ``UnitStatement`` (unit_movement_service) as a branded PDF, in the design of the
wallet's account statement, or as an Excel workbook (two sheets: movements, and units by
property). Both only lay out what the service computed."""

from __future__ import annotations

import datetime as dt
import io

from reportlab.pdfgen import canvas

from app.services.certificate_service import (
    _CX,
    _GOLD,
    _GREEN,
    _GREEN_D,
    _H,
    _INK,
    _MUTED,
    _PANEL,
    _W,
    _clip,
    _divider,
    _draw_brand,
    _fit_centred,
    _wrap,
)
from app.services.installment_pdf import _RED, _page_frame
from app.services.statement_render import _spaced_centred
from app.services.unit_movement_service import UnitStatement

_BOTTOM = 112.0
_ROW_H = 15.0
_C_DATE = 50.0
_C_PROP = 100.0
_C_WHAT = 222.0
_C_WHO = 394.0
_C_UNITS = 505.0  # right-aligned
_C_HELD = 548.0  # right-aligned


def _day(d: dt.date | dt.datetime) -> str:
    return d.strftime("%d %b %Y")


def _signed(n: int) -> str:
    if n == 0:
        return "0"
    return f"+{n:,}" if n > 0 else f"-{abs(n):,}"


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def _footer(c: canvas.Canvas, stmt: UnitStatement, page: int) -> None:
    c.setFillColor(_MUTED)
    c.setFont("Times-Roman", 8.5)
    note = (
        "From the Capimax PropShare ownership ledger; times are UTC, both dates included. "
        "HELD: units held in that property after the movement. Informational only: not a "
        "tax document or a title deed."
    )
    fy = 96
    for ln in _wrap(note, "Times-Roman", 8.5, 400):
        c.drawCentredString(_CX, fy, ln)
        fy -= 11
    c.setFillColor(_GOLD)
    c.setFont("Helvetica", 8)
    c.drawCentredString(_CX, 62, f"{stmt.ref}   •   page {page}   •   capimaxpropshare.com")


def _movements_header(c: canvas.Canvas, y: float) -> float:
    c.setFillColor(_MUTED)
    c.setFont("Helvetica-Bold", 7.5)
    c.drawString(_C_DATE, y, "DATE")
    c.drawString(_C_PROP, y, "PROPERTY")
    c.drawString(_C_WHAT, y, "MOVEMENT")
    c.drawString(_C_WHO, y, "OTHER PARTY")
    c.drawRightString(_C_UNITS, y, "UNITS")
    c.drawRightString(_C_HELD, y, "HELD")
    c.setStrokeColor(_GOLD)
    c.setLineWidth(0.8)
    c.line(_C_DATE, y - 6, _C_HELD, y - 6)
    return y - 19


def render_pdf(stmt: UnitStatement) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(_W, _H), pageCompression=0)
    c.setTitle("Capimax PropShare — Unit Statement")
    page = 1
    _page_frame(c)

    def new_page(title: str) -> float:
        nonlocal page
        _footer(c, stmt, page)
        c.showPage()
        page += 1
        _page_frame(c)
        _spaced_centred(c, _CX, 780, "Times-Bold", 15, title, 1.0, _GREEN_D)
        return 748.0

    _draw_brand(c, _CX)
    _divider(c, _CX, 712, 150)
    _spaced_centred(c, _CX, 676, "Times-Bold", 27, "Unit Statement", 1.0, _GREEN_D)
    c.setFillColor(_MUTED)
    c.setFont("Times-Italic", 11.5)
    c.drawCentredString(_CX, 658, f"{_day(stmt.start)} – {_day(stmt.end)}  (UTC)")
    _fit_centred(c, _CX, 630, "Times-Bold", 15, stmt.holder, 440, _GREEN)
    c.setFillColor(_INK)
    c.setFont("Times-Roman", 10.5)
    c.drawCentredString(_CX, 613, f"{stmt.email}   •   {stmt.ref}")

    # summary panel: 3 columns x 2 rows
    px, pw, ph, py = 56.0, _W - 112.0, 84.0, 510.0
    c.setFillColor(_PANEL)
    c.setStrokeColor(_GOLD)
    c.setLineWidth(0.8)
    c.roundRect(px, py, pw, ph, 6, stroke=1, fill=1)
    held_start = sum(p.opening for p in stmt.properties)
    held_end = sum(p.closing for p in stmt.properties)
    fields = [
        ("UNITS HELD AT THE START", f"{held_start:,}"),
        ("UNITS IN", f"{stmt.units_in:,}"),
        ("UNITS OUT", f"{stmt.units_out:,}"),
        ("UNITS HELD AT THE END", f"{held_end:,}"),
        ("MOVEMENTS", str(len(stmt.rows))),
        ("ISSUED", f"{_day(stmt.generated_at)} {stmt.generated_at:%H:%M} UTC"),
    ]
    col_w = pw / 3
    for i, (label, val) in enumerate(fields):
        col, row = i % 3, i // 3
        x = px + 18 + col * col_w
        y = py + ph - 24 - row * 34
        c.setFillColor(_MUTED)
        c.setFont("Helvetica", 6.8)
        c.drawString(x, y, label)
        strong = label == "UNITS HELD AT THE END"
        c.setFillColor(_GREEN_D if strong else _INK)
        c.setFont("Times-Bold" if strong else "Times-Roman", 10.5)
        c.drawString(x, y - 15, _clip(val, "Times-Roman", 10.5, col_w - 24))

    # movements
    y = py - 26
    c.setFillColor(_GREEN_D)
    c.setFont("Times-Bold", 12)
    c.drawString(_C_DATE, y, "Unit movements")
    y = _movements_header(c, y - 20)
    if not stmt.rows:
        c.setFillColor(_MUTED)
        c.setFont("Times-Italic", 10)
        c.drawString(_C_DATE, y, "No unit movements in this period.")
        y -= _ROW_H
    for i, r in enumerate(stmt.rows):
        if y < _BOTTOM:
            y = _movements_header(c, new_page("Unit movements (continued)"))
        if i % 2:
            c.setFillColor(_PANEL)
            c.rect(_C_DATE - 6, y - 4, _C_HELD - _C_DATE + 12, _ROW_H, stroke=0, fill=1)
        c.setFillColor(_MUTED)
        c.setFont("Times-Roman", 8.5)
        c.drawString(_C_DATE, y, _day(r.at))
        c.setFillColor(_INK)
        c.drawString(_C_PROP, y, _clip(r.property_title, "Times-Roman", 8.5, _C_WHAT - _C_PROP - 8))
        c.drawString(_C_WHAT, y, _clip(r.label, "Times-Roman", 8.5, _C_WHO - _C_WHAT - 8))
        c.setFillColor(_MUTED)
        c.drawString(
            _C_WHO, y, _clip(r.counterparty or "-", "Times-Roman", 8.5, _C_UNITS - _C_WHO - 34)
        )
        c.setFillColor(_GREEN if r.units >= 0 else _RED)
        c.setFont("Times-Bold", 8.5)
        c.drawRightString(_C_UNITS, y, _signed(r.units))
        c.setFillColor(_INK)
        c.setFont("Times-Roman", 8.5)
        c.drawRightString(_C_HELD, y, f"{r.balance:,}")
        y -= _ROW_H

    # units by property
    title = "Units by property"
    if y < _BOTTOM + 60:  # title + header + one row must fit, else start a page
        y = new_page(title)
    else:
        y -= 18
        c.setFillColor(_GREEN_D)
        c.setFont("Times-Bold", 12)
        c.drawString(_C_DATE, y, title)
        y -= 20
    cols = (("AT THE START", 352.0), ("IN", 418.0), ("OUT", 484.0), ("AT THE END", _C_HELD))
    c.setFillColor(_MUTED)
    c.setFont("Helvetica-Bold", 7.5)
    c.drawString(_C_DATE, y, "PROPERTY")
    for label, x in cols:
        c.drawRightString(x, y, label)
    c.setStrokeColor(_GOLD)
    c.line(_C_DATE, y - 6, _C_HELD, y - 6)
    y -= 19
    if not stmt.properties:
        c.setFillColor(_MUTED)
        c.setFont("Times-Italic", 10)
        c.drawString(_C_DATE, y, "No units held in this period.")
    for p in stmt.properties:
        if y < _BOTTOM:
            y = new_page(f"{title} (continued)")
        c.setFillColor(_INK)
        c.setFont("Times-Roman", 9)
        c.drawString(_C_DATE, y, _clip(p.property_title, "Times-Roman", 9, 352.0 - _C_DATE - 70))
        for (_label, x), v in zip(
            cols, (f"{p.opening:,}", _signed(p.units_in), _signed(-p.units_out)), strict=False
        ):
            c.drawRightString(x, y, v)
        c.setFont("Times-Bold", 9)
        c.setFillColor(_GREEN_D)
        c.drawRightString(_C_HELD, y, f"{p.closing:,}")
        y -= _ROW_H

    _footer(c, stmt, page)
    c.showPage()
    c.save()
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def render_xlsx(stmt: UnitStatement) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="0F6E4C")
    head_font = Font(bold=True, color="FFFFFF")
    signed = "+#,##0;[Red]-#,##0;0"

    wb = Workbook()
    ws = wb.active
    ws.title = "Unit movements"
    meta = [
        ("Capimax PropShare — Unit statement", None),
        ("Holder", stmt.holder),
        ("Email", stmt.email),
        ("Reference", stmt.ref),
        ("Period (UTC, inclusive)", f"{stmt.start.isoformat()} to {stmt.end.isoformat()}"),
        ("Units held at the start", sum(p.opening for p in stmt.properties)),
        ("Units in", stmt.units_in),
        ("Units out", stmt.units_out),
        ("Units held at the end", sum(p.closing for p in stmt.properties)),
        ("Issued (UTC)", stmt.generated_at.strftime("%Y-%m-%d %H:%M")),
    ]
    for label, value in meta:
        ws.append([label, value])
    ws["A1"].font = Font(bold=True, size=14, color="0F6E4C")
    for row in ws.iter_rows(min_row=2, max_row=len(meta), max_col=1):
        row[0].font = bold
    ws.append([])

    header = [
        "Date (UTC)",
        "Property",
        "Movement",
        "Other party",
        "Units",
        "Held after",
        "Reference",
    ]
    ws.append(header)
    head_row = ws.max_row
    for cell in ws[head_row]:
        cell.font, cell.fill = head_font, head_fill
    for r in stmt.rows:
        ws.append(
            [
                r.at.astimezone(dt.UTC).replace(tzinfo=None),
                r.property_title,
                r.label,
                r.counterparty,
                r.units,
                r.balance,
                r.reference,
            ]
        )
        ws.cell(row=ws.max_row, column=1).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(row=ws.max_row, column=5).number_format = signed
    if not stmt.rows:
        ws.append(["No unit movements in this period."])
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    for col, width in zip("ABCDEFG", (24, 38, 40, 28, 10, 12, 12), strict=True):
        ws.column_dimensions[col].width = width

    props = wb.create_sheet("Units by property")
    props.append(["Property", "At the start", "In", "Out", "At the end"])
    for cell in props[1]:
        cell.font, cell.fill = head_font, head_fill
    for p in stmt.properties:
        props.append([p.property_title, p.opening, p.units_in, p.units_out, p.closing])
    props.append(
        [
            "Total",
            sum(p.opening for p in stmt.properties),
            stmt.units_in,
            stmt.units_out,
            sum(p.closing for p in stmt.properties),
        ]
    )
    for cell in props[props.max_row]:
        cell.font = bold
    for col, width in zip("ABCDE", (44, 14, 10, 10, 14), strict=True):
        props.column_dimensions[col].width = width
    for sheet in (ws, props):
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
