"""Account statements (PDF / Excel for a period the user picks).

What each test protects:
  * the arithmetic: opening = everything booked before the period, closing = opening + the
    period's signed movements, running balance on every row, money in / out split by type;
    rows just outside the period (by one day) are excluded, both end dates are included;
  * only the signed-in user's own ledger is read; holdings are the units owned at the END
    of the period (a later purchase is not shown);
  * the PDF carries the real figures and paginates long periods; the Excel file has the three
    sheets with numeric cells (not text), so a spreadsheet can sum them;
  * bad periods are refused with a clear code; the export is audited and rate-limited.
"""

# ruff: noqa: E501
from __future__ import annotations

import datetime as dt
import decimal
import io
import re
import uuid

import pytest

PW = "Passw0rd!23"
D = decimal.Decimal


async def _user(client, db, email: str) -> tuple[str, uuid.UUID]:
    r = await client.post(
        "/api/v1/auth/register", json={"email": email, "password": PW, "full_name": "Sara Investor"}
    )
    assert r.status_code == 201, r.text
    return r.json()["access_token"], db("SELECT id FROM users WHERE email=:e", e=email)[0][0]


def _tx(db, uid, when: str, kind: str, amount: str, status="completed", desc=None):
    db(
        "INSERT INTO transactions (user_id, type, amount, status, description, created_at) "
        "VALUES (:u, CAST(:t AS transaction_type), :a, :s, :d, CAST(:w AS timestamptz))",
        u=uid,
        t=kind,
        a=amount,
        s=status,
        d=desc,
        w=when,
    )


def _prop(db, title: str) -> str:
    pid = str(uuid.uuid4())
    db(
        "INSERT INTO properties (id,title,slug,location,city,country,property_type,model,status,total_value,unit_price,"
        "total_units,available_units,minimum_investment,description) VALUES "
        "(:id,:t,:s,'Dubai','Dubai','UAE','apartment','ready-income','active',100000,100,1000,1000,100,'x')",
        id=pid,
        t=title,
        s=f"p-{pid[:8]}",
    )
    return pid


def _own(db, uid, pid, units: int, when: str):
    db(
        "INSERT INTO ownership_ledger (user_id, property_id, units, unit_price, reason, created_at) "
        "VALUES (:u,:p,:n,100,'purchase', CAST(:w AS timestamptz))",
        u=uid,
        p=pid,
        n=units,
        w=when,
    )


def _h(tok: str) -> dict:
    return {"Authorization": f"Bearer {tok}"}


def _seed_history(db, uid):
    _tx(db, uid, "2026-01-05T10:00:00Z", "deposit", "1000.00", desc="Card deposit")
    _tx(db, uid, "2026-01-31T23:59:59Z", "fee", "-1.00", desc="Before the period")
    # --- period 2026-02-01 .. 2026-03-02 ---
    _tx(db, uid, "2026-02-01T00:00:00Z", "investment", "-300.00", desc="Creek Tower, 3 units")
    _tx(db, uid, "2026-02-01T00:00:00Z", "fee", "-7.50", desc="Platform fee")
    _tx(db, uid, "2026-03-01T08:00:00Z", "return", "12.25", desc="March distribution")
    _tx(
        db,
        uid,
        "2026-03-02T23:59:00Z",
        "withdrawal",
        "-100.00",
        status="pending",
        desc="Withdrawal hold",
    )
    # --- after ---
    _tx(
        db,
        uid,
        "2026-03-03T00:00:00Z",
        "withdrawal",
        "100.00",
        status="reversed",
        desc="Withdrawal reversed",
    )
    _balance(db, uid, "703.75")  # == SUM(ledger): the wallet invariant holds


def _balance(db, uid, amount: str):
    db("UPDATE wallets SET balance=:b WHERE user_id=:u", b=amount, u=uid)


@pytest.mark.asyncio
async def test_statement_arithmetic_and_scope(client, db, asession):
    from app.services import statement_service

    _tok, uid = await _user(client, db, "stmt@x.io")
    _other_tok, other = await _user(client, db, "other@x.io")
    _seed_history(db, uid)
    _tx(db, other, "2026-02-10T00:00:00Z", "deposit", "5000.00")  # never in sara's statement
    _balance(db, other, "5000.00")
    tower, marina = _prop(db, "Creek Tower"), _prop(db, "Marina Heights")
    _own(db, uid, tower, 3, "2026-02-01T00:00:00Z")
    _own(db, uid, marina, 5, "2026-03-10T00:00:00Z")  # bought after the period

    s = await statement_service.build_statement(
        asession, uid, dt.date(2026, 2, 1), dt.date(2026, 3, 2)
    )
    assert s.opening == D("999.00")
    assert [r.amount for r in s.rows] == [D("-300.00"), D("-7.50"), D("12.25"), D("-100.00")]
    assert [r.balance for r in s.rows] == [D("699.00"), D("691.50"), D("703.75"), D("603.75")]
    assert s.closing == D("603.75")
    assert (s.money_in, s.money_out) == (D("12.25"), D("407.50"))
    assert s.rows[3].status == "In progress" and s.rows[2].type_label == "Return / distribution"
    totals = {t.label: (t.money_in, t.money_out) for t in s.totals}
    assert totals["Fee"] == (D("0.00"), D("7.50"))
    assert totals["Investment"] == (D("0.00"), D("300.00"))
    assert [(h.property_title, h.units) for h in s.holdings] == [("Creek Tower", 3)]
    assert s.ref.startswith("STM-") and s.ref.endswith("20260201-20260302")

    # an empty period still has correct opening == closing
    empty = await statement_service.build_statement(
        asession, uid, dt.date(2026, 4, 1), dt.date(2026, 4, 30)
    )
    assert empty.rows == [] and empty.opening == empty.closing == D("703.75")


@pytest.mark.asyncio
async def test_statement_ends_on_the_real_balance_when_history_is_incomplete(client, db, asession):
    """A balance set by hand (demo data) is not in the ledger: the statement must still end on
    the balance the user sees, never on an impossible negative figure."""
    from app.services import statement_service

    _tok, uid = await _user(client, db, "demo@x.io")
    _tx(db, uid, "2026-09-20T10:00:00Z", "investment", "-1000.00")
    _balance(db, uid, "4000.00")  # the 5,000 top-up was never booked in the ledger
    s = await statement_service.build_statement(
        asession, uid, dt.date(2026, 9, 1), dt.date(2026, 9, 21)
    )
    assert (s.opening, s.closing) == (D("5000.00"), D("4000.00"))
    assert s.rows[0].balance == D("4000.00")


@pytest.mark.asyncio
async def test_pdf_and_excel_downloads(client, db):
    from openpyxl import load_workbook

    tok, uid = await _user(client, db, "dl@x.io")
    _seed_history(db, uid)
    q = "start=2026-02-01&end=2026-03-02"

    r = await client.get(f"/api/v1/wallet/statement?{q}&format=pdf", headers=_h(tok))
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/pdf"
    assert r.headers["content-disposition"] == (
        'attachment; filename="capimax-statement-2026-02-01-to-2026-03-02.pdf"'
    )
    assert r.headers["cache-control"] == "no-store"
    pdf = r.content
    assert pdf.startswith(b"%PDF")
    for needle in (
        b"Account Statement",
        b"603.75 USD",
        b"999.00 USD",
        b"+12.25",
        b"-300.00",
        b"Sara Investor",
    ):
        assert needle in pdf, needle

    r = await client.get(f"/api/v1/wallet/statement?{q}&format=xlsx", headers=_h(tok))
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Statement", "Summary by type", "Holdings"]
    ws = wb["Statement"]
    meta = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=2).value for i in range(1, 12)}
    assert meta["Opening balance"] == 999 and meta["Closing balance"] == pytest.approx(603.75)
    header_row = next(i for i in range(1, 20) if ws.cell(row=i, column=1).value == "Date (UTC)")
    amounts = [ws.cell(row=header_row + k, column=5).value for k in range(1, 5)]
    assert amounts == [-300, -7.5, 12.25, -100]  # numbers, not text
    assert isinstance(ws.cell(row=header_row + 1, column=1).value, dt.datetime)
    summary = wb["Summary by type"]
    assert summary.cell(row=summary.max_row, column=1).value == "Total"
    assert summary.cell(row=summary.max_row, column=3).value == pytest.approx(407.5)

    assert db("SELECT count(*) FROM audit_log WHERE action='statement.exported'")[0][0] == 2


@pytest.mark.asyncio
async def test_statement_is_english_only(client, db, asession):
    """Owner's decision (2026-09-24): the statement is in English. Text in another script (an
    Arabic name, property title or admin note) is replaced by an English equivalent instead of
    being printed as empty boxes; English text keeps what the fonts can draw."""
    from pypdf import PdfReader

    from app.services import statement_service as svc

    tok, uid = await _user(client, db, "arabic@x.io")
    db("UPDATE users SET full_name='سارة المستثمرة' WHERE id=:u", u=uid)
    pid = _prop(db, "برج الخور")
    db("UPDATE properties SET slug='creek-tower' WHERE id=:p", p=pid)
    _own(db, uid, pid, 2, "2026-02-10T00:00:00Z")
    _tx(db, uid, "2026-02-10T00:00:00Z", "deposit", "500.00", desc="Deposit ✓ via card")
    _tx(db, uid, "2026-02-11T00:00:00Z", "withdrawal", "-50.00", desc="سحب — الحساب غير صحيح")
    _tx(db, uid, "2026-02-12T00:00:00Z", "fee", "-1.00", desc="Fee ref ١٢٣")
    _tx(db, uid, "2026-02-13T00:00:00Z", "return", "3.00")
    _balance(db, uid, "452.00")

    s = await svc.build_statement(asession, uid, dt.date(2026, 2, 1), dt.date(2026, 2, 28))
    assert s.holder == "arabic@x.io"
    assert [r.description for r in s.rows] == ["Deposit via card", "Withdrawal", "Fee", ""]
    assert [(h.property_title, h.units) for h in s.holdings] == [("Creek Tower", 2)]
    for text in (
        s.holder,
        *(r.description for r in s.rows),
        *(h.property_title for h in s.holdings),
    ):
        text.encode("cp1252")  # everything is drawable by the PDF's standard fonts

    r = await client.get(
        "/api/v1/wallet/statement?start=2026-02-01&end=2026-02-28&format=pdf", headers=_h(tok)
    )
    text = "".join(page.extract_text() for page in PdfReader(io.BytesIO(r.content)).pages)
    assert "arabic@x.io" in text and "Creek Tower" in text and "Deposit via card" in text
    assert not re.search("[؀-ۿ]", text)


@pytest.mark.asyncio
async def test_long_statement_paginates(client, db):
    tok, uid = await _user(client, db, "long@x.io")
    for i in range(120):
        _tx(db, uid, f"2026-05-{1 + i % 28:02d}T{i % 24:02d}:00:00Z", "deposit", "10.00")
    _balance(db, uid, "1200.00")
    r = await client.get(
        "/api/v1/wallet/statement?start=2026-05-01&end=2026-05-31&format=pdf", headers=_h(tok)
    )
    assert r.status_code == 200
    pages = len(re.findall(rb"/Type /Page[^s]", r.content))
    assert pages == 4 and b"page 4" in r.content  # 120 rows: 1 first page + 3 continued
    assert b"1,200.00 USD" in r.content


@pytest.mark.parametrize(
    ("query", "code"),
    [
        ("start=2026-03-02&end=2026-02-01", "BAD_PERIOD"),
        ("start=2026-01-01&end=2999-01-01", "BAD_PERIOD"),
        ("start=2019-12-31&end=2020-01-31", "BAD_PERIOD"),
        ("start=2020-01-01&end=2024-01-01", "PERIOD_TOO_LONG"),
    ],
)
@pytest.mark.asyncio
async def test_bad_periods_are_refused(client, db, query, code):
    tok, _uid = await _user(client, db, "bad@x.io")
    r = await client.get(f"/api/v1/wallet/statement?{query}", headers=_h(tok))
    assert r.status_code == 422
    assert r.json()["error"]["code"] == code


@pytest.mark.asyncio
async def test_statement_needs_sign_in_and_a_known_format(client, db):
    assert (
        await client.get("/api/v1/wallet/statement?start=2026-01-01&end=2026-01-31")
    ).status_code == 401
    tok, _uid = await _user(client, db, "fmt@x.io")
    r = await client.get(
        "/api/v1/wallet/statement?start=2026-01-01&end=2026-01-31&format=csv", headers=_h(tok)
    )
    assert r.status_code == 422
