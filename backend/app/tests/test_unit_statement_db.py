"""The unit statement (client, 2026-10-15): the wallet has a statement of its money, the units
had none. A member downloads, for a period they choose, every unit movement with the units
held afterwards, and per property the units held at the start and at the end, as a PDF or an
Excel workbook. Only their own ledger lines are read.
"""

from __future__ import annotations

import datetime as dt
import io
import uuid

import pytest

PW = "Passw0rd!23"


async def _user(client, db, email: str, name: str) -> tuple[str, dict]:
    r = await client.post(
        "/api/v1/auth/register", json={"email": email, "password": PW, "full_name": name}
    )
    assert r.status_code == 201, r.text
    uid = str(db("SELECT id FROM users WHERE email=:e", e=email)[0][0])
    return uid, {"Authorization": f"Bearer {r.json()['access_token']}"}


def _property(db, title: str) -> str:
    pid = str(uuid.uuid4())
    db(
        "INSERT INTO properties (id,title,location,property_type,model,status,total_value,"
        "unit_price,total_units,available_units,minimum_investment) VALUES "
        "(:id,:t,'Dubai','residential','ready-income','active',100000,100,1000,1000,100)",
        id=pid,
        t=title,
    )
    return pid


def _ledger(db, uid: str, pid: str, units: int, reason: str, at: dt.datetime, ref=None) -> None:
    db(
        "INSERT INTO ownership_ledger (user_id, property_id, units, unit_price, reason, "
        "created_at, ref_id) VALUES (:u,:p,:n,100,:r,:at,:ref)",
        u=uid,
        p=pid,
        n=units,
        r=reason,
        at=at,
        ref=ref,
    )


def _at(day: str) -> dt.datetime:
    return dt.datetime.fromisoformat(day + "T10:00:00+00:00")


@pytest.mark.asyncio
async def test_the_unit_statement_lists_the_period_with_what_was_held(client, db):
    me, h = await _user(client, db, "ahmed@units.io", "Ahmed Mansouri")
    other, _oh = await _user(client, db, "someone@units.io", "Someone Else")
    marina, palm = _property(db, "Marina Heights"), _property(db, "Palm Court")
    # before the period: 10 units of Marina
    _ledger(db, me, marina, 10, "purchase", _at("2026-08-20"))
    # in the period
    _ledger(db, me, marina, 5, "secondary_purchase", _at("2026-09-02"))
    _ledger(db, me, palm, 8, "purchase", _at("2026-09-05"))
    _ledger(db, me, marina, -3, "secondary_sale", _at("2026-09-10"))
    # after the period, and someone else's: not in it
    _ledger(db, me, palm, 2, "purchase", _at("2026-10-05"))
    _ledger(db, other, marina, 50, "purchase", _at("2026-09-03"))

    q = "/api/v1/investments/movements/statement?start=2026-09-01&end=2026-09-30"
    r = await client.get(q + "&format=xlsx", headers=h)
    assert r.status_code == 200, r.text
    assert (
        "capimax-unit-statement-2026-09-01-to-2026-09-30.xlsx" in r.headers["content-disposition"]
    )
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Unit movements"]
    meta = {row[0]: row[1] for row in ws.iter_rows(min_row=1, max_row=10, values_only=True)}
    assert meta["Holder"] == "Ahmed Mansouri"
    assert meta["Units held at the start"] == 10
    assert (meta["Units in"], meta["Units out"], meta["Units held at the end"]) == (13, 3, 20)
    rows = [
        row for row in ws.iter_rows(min_row=13, values_only=True) if isinstance(row[0], dt.datetime)
    ]
    assert [(r[1], r[2], r[4], r[5]) for r in rows] == [
        ("Marina Heights", "Bought on the secondary market", 5, 15),
        ("Palm Court", "Purchase", 8, 8),
        ("Marina Heights", "Sold on the secondary market", -3, 12),
    ]
    props = [row for row in wb["Units by property"].iter_rows(min_row=2, values_only=True)]
    assert props == [
        ("Marina Heights", 10, 5, 3, 12),
        ("Palm Court", 0, 8, 0, 8),
        ("Total", 10, 13, 3, 20),
    ]

    r = await client.get(q, headers=h)
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF")
    for text in (b"Marina Heights", b"Sold on the secondary market", b"Units by property"):
        assert text in r.content, text
    assert db("SELECT count(*) FROM audit_log WHERE action='unit_statement.exported'")[0][0] == 2


@pytest.mark.asyncio
async def test_a_transfer_names_the_other_person_in_the_statement(client, db):
    owner, h = await _user(client, db, "ahmed2@units.io", "Ahmed Mansouri")
    son, sh = await _user(client, db, "omar@units.io", "Omar Mansouri")
    pid = _property(db, "Marina Heights")
    group = str(uuid.uuid4())
    db("INSERT INTO family_groups (id, owner_id, name) VALUES (:g,:o,'Family')", g=group, o=owner)
    m_owner, m_son = str(uuid.uuid4()), str(uuid.uuid4())
    db(
        "INSERT INTO family_members (id, family_group_id, user_id, name, relationship, "
        "is_verified) VALUES (:a,:g,:o,'Ahmed','Self (Owner)',true), "
        "(:b,:g,:s,'Omar','Son',true)",
        a=m_owner,
        b=m_son,
        g=group,
        o=owner,
        s=son,
    )
    transfer = str(uuid.uuid4())
    db(
        "INSERT INTO family_transfers (id, family_group_id, from_member_id, to_member_id, "
        "property_id, units, status) VALUES (:t,:g,:a,:b,:p,4,'completed')",
        t=transfer,
        g=group,
        a=m_owner,
        b=m_son,
        p=pid,
    )
    _ledger(db, owner, pid, 10, "purchase", _at("2026-09-01"))
    _ledger(db, owner, pid, -4, "family_transfer_out", _at("2026-09-02"), ref=transfer)
    _ledger(db, son, pid, 4, "family_transfer_in", _at("2026-09-02"), ref=transfer)

    from openpyxl import load_workbook

    q = "/api/v1/investments/movements/statement?start=2026-09-01&end=2026-09-30&format=xlsx"
    mine = load_workbook(io.BytesIO((await client.get(q, headers=h)).content))
    theirs = load_workbook(io.BytesIO((await client.get(q, headers=sh)).content))

    def lines(wb):
        return [
            (r[2], r[3], r[4])
            for r in wb["Unit movements"].iter_rows(min_row=13, values_only=True)
            if isinstance(r[0], dt.datetime)
        ]

    assert lines(mine) == [("Purchase", None, 10), ("Transferred to a family member", "Omar", -4)]
    assert lines(theirs) == [("Received from a family member", "Ahmed Mansouri", 4)]


@pytest.mark.asyncio
async def test_the_unit_statement_is_the_members_own_and_checks_its_period(client, db):
    _me, h = await _user(client, db, "ahmed3@units.io", "Ahmed Mansouri")
    base = "/api/v1/investments/movements/statement"
    assert (await client.get(f"{base}?start=2026-09-01&end=2026-09-30")).status_code == 401
    r = await client.get(f"{base}?start=2026-09-30&end=2026-09-01", headers=h)
    assert r.status_code == 422 and r.json()["error"]["code"] == "BAD_PERIOD"
    future = (dt.datetime.now(dt.UTC).date() + dt.timedelta(days=3)).isoformat()
    r = await client.get(f"{base}?start=2026-09-01&end={future}", headers=h)
    assert r.status_code == 422
    # nothing held and nothing moved: an honest, empty statement
    r = await client.get(f"{base}?start=2026-09-01&end=2026-09-30", headers=h)
    assert r.status_code == 200 and b"No unit movements in this period." in r.content
